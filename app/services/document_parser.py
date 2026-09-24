from pathlib import Path
import re

from pypdf import PdfReader
from docx import Document
import pandas as pd
import pdfplumber
import pytesseract
import pymupdf
from PIL import Image, ImageEnhance, ImageFilter, ImageOps
from io import BytesIO


def classify_ocr_token(word):
    """
    Classify an OCR token so important structured values
    are protected from automatic spelling correction.
    """

    # Date formats such as:
    # 31-May-2019
    # 31/05/2019
    # 31-05-2019
    if re.fullmatch(
        r"\d{1,2}[-/](?:[A-Za-z]{3,9}|\d{1,2})[-/]\d{2,4}",
        word
    ):
        return "date"

    # Numbers such as:
    # 3329
    # 2019
    # 12
    if re.fullmatch(r"\d+", word):
        return "number"

    # Alphanumeric IDs such as:
    # ABC12345
    # 2T3INE0641H
    if re.fullmatch(r"[A-Za-z0-9]+", word) and any(
        char.isdigit() for char in word
    ):
        return "id_or_code"

    # Normal alphabetic words
    if re.search(r"[A-Za-z]", word):
        return "text"

    return "other"


def preprocess_ocr_image(image):
    """
    Improve a scanned document image before OCR.
    """


def preprocess_ocr_image(image):
    """
    Improve a scanned document image before OCR.
    """

    # 1. Convert to grayscale
    image = ImageOps.grayscale(image)

    # 2. Increase contrast automatically
    image = ImageOps.autocontrast(image)

    # 3. Increase image size
    width, height = image.size
    image = image.resize(
        (width * 2, height * 2),
        Image.Resampling.LANCZOS
    )

    # 4. Reduce small image noise
    image = image.filter(ImageFilter.MedianFilter(size=3))

    # 5. Slightly sharpen text
    image = image.filter(ImageFilter.SHARPEN)

    # 6. Enhance contrast again
    contrast = ImageEnhance.Contrast(image)
    image = contrast.enhance(1.5)

    return image

def parse_pdf(file_path):
    """Extract text and tables from a PDF, using OCR when necessary."""

    content = []

    with pdfplumber.open(file_path) as pdf:

        for page_number, page in enumerate(pdf.pages, start=1):

            # Try normal text extraction first
            page_text = page.extract_text()

            # If normal extraction finds text
            if page_text and page_text.strip():

                content.append({
                    "type": "text",
                    "page": page_number,
                    "content": page_text
                })

            # If no text was found, use OCR
            else:

                print(f"Page {page_number}: No text found. Running OCR...")

                pdf_document = pymupdf.open(file_path)
                pdf_page = pdf_document[page_number - 1]

                # Render PDF page as an image
                pix = pdf_page.get_pixmap(matrix=pymupdf.Matrix(3, 3))

                # Convert image to bytes
                image_bytes = pix.tobytes("png")

                image = Image.open(BytesIO(image_bytes))

                # Preprocess image before OCR
                processed_image = preprocess_ocr_image(image)

                # Run OCR
                ocr_data = pytesseract.image_to_data(
                    processed_image,
                    config="--psm 6",
                    output_type=pytesseract.Output.DICT
                )

                ocr_words = []
                confidence_values = []
                low_confidence_words = []

                for i in range(len(ocr_data["text"])):

                    word = ocr_data["text"][i].strip()
                    confidence = float(ocr_data["conf"][i])

                    if word:
                        ocr_words.append(word)

                        if confidence >= 0:
                            confidence_values.append(confidence)

                            # Track meaningful low-confidence words
                            if confidence < 60:

                                if len(word) >= 3 and re.search(r"[A-Za-z]", word):

                                    token_type = classify_ocr_token(word)

                                    low_confidence_words.append({
                                        "word": word,
                                        "confidence": round(confidence, 2),
                                        "type": token_type
                                    })
                if confidence_values:
                    average_confidence = sum(confidence_values) / len(confidence_values)
                else:
                    average_confidence = 0

                # Determine OCR confidence status
                if average_confidence >= 85:
                    confidence_status = "HIGH"
                elif average_confidence >= 60:
                    confidence_status = "MEDIUM"
                else:
                    confidence_status = "LOW"

                ocr_text = " ".join(ocr_words)

                if ocr_text.strip():

                    content.append({
                        "type": "ocr_text",
                        "page": page_number,
                        "content": ocr_text,
                        "average_confidence": round(average_confidence, 2),
                        "confidence_status": confidence_status,
                        "low_confidence_words": low_confidence_words
                    })

                pdf_document.close()

            # Extract tables
            tables = page.extract_tables()

            for table in tables:

                if table:

                    cleaned_rows = []

                    for row in table:

                        cleaned_row = [
                            cell.strip() if cell else ""
                            for cell in row
                        ]

                        cleaned_rows.append(cleaned_row)

                    content.append({
                        "type": "table",
                        "page": page_number,
                        "rows": cleaned_rows
                    })

    return content


def parse_docx(file_path):
    """Extract structured content from a DOCX file."""

    document = Document(file_path)

    content = []

    # Extract paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        # Detect basic paragraph style
        style = paragraph.style.name.lower()

        if "heading" in style:
            content.append({
                "type": "heading",
                "text": text
            })

        elif "list" in style:
            content.append({
                "type": "bullet",
                "text": text
            })

        else:
            content.append({
                "type": "paragraph",
                "text": text
            })

    # Extract tables
    for table in document.tables:

        rows = []

        for row in table.rows:

            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            rows.append(row_data)

        content.append({
            "type": "table",
            "rows": rows
        })

    return content


def parse_txt(file_path):
    """Extract text from a TXT file."""

    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()


def parse_xlsx(file_path):
    """Extract data from an Excel file."""

    excel_file = pd.ExcelFile(file_path)

    content = []

    for sheet_name in excel_file.sheet_names:

        dataframe = pd.read_excel(
            file_path,
            sheet_name=sheet_name
        )

        content.append(f"--- Sheet: {sheet_name} ---")

        content.append(
            dataframe.to_string(index=False)
        )

    return "\n".join(content)


def parse_document(file_path):
    """Detect the file type and extract its content."""

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return parse_pdf(file_path)

    elif extension == ".docx":
        return parse_docx(file_path)

    elif extension == ".txt":
        return parse_txt(file_path)

    elif extension in [".xlsx", ".xls"]:
        return parse_xlsx(file_path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )