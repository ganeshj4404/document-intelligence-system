from pptx import Presentation


presentation = Presentation()

# Slide 1
slide = presentation.slides.add_slide(
    presentation.slide_layouts[0]
)

slide.shapes.title.text = "Document Intelligence Test"

slide.placeholders[1].text = (
    "This is a test PowerPoint document."
)

# Slide 2
slide = presentation.slides.add_slide(
    presentation.slide_layouts[1]
)

slide.shapes.title.text = "Project Details"

slide.placeholders[1].text = (
    "Parser testing\n"
    "LLM summarization\n"
    "Text-to-Speech"
)

# Save
presentation.save(
    "data/input/test.pptx"
)

print("Test PowerPoint created successfully.")